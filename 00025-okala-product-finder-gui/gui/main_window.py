import os

try:
    import openpyxl
except ImportError:
    openpyxl = None  # We'll handle missing module

from PyQt5.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QFileDialog, QTextEdit, QGroupBox, QPlainTextEdit,
    QMessageBox, QTreeWidget, QTreeWidgetItem, QProgressBar
)
from PyQt5.QtCore import Qt, pyqtSlot
from gui.config import load_settings, save_settings, load_store_locations
from gui.workers import StoreUpdateWorker, ProductIngestionWorker


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Appu Okala Product Finder")
        self.resize(1300, 850)

        # ---- State ----
        self.settings = load_settings()
        self.store_details = []
        self.product_ids_for_search = []
        self.store_worker = None
        self.product_worker = None

        # ---- Build UI ----
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QHBoxLayout(central_widget)
        main_layout.setContentsMargins(8, 8, 8, 8)

        # ===== LEFT PANEL =====
        left_panel = QWidget()
        left_layout = QVBoxLayout(left_panel)
        left_layout.setContentsMargins(0, 0, 5, 0)

        # 1) Auth Token group
        auth_group = QGroupBox("Authentication")
        auth_layout = QVBoxLayout(auth_group)

        self.auth_status_label = QLabel()
        self.update_auth_status()
        auth_layout.addWidget(self.auth_status_label)

        self.token_edit = QPlainTextEdit()
        self.token_edit.setPlaceholderText("Paste your Bearer token here...")
        self.token_edit.setMaximumBlockCount(3)
        self.token_edit.setFixedHeight(80)
        token_value = self.settings.get("auth_token", "")
        if token_value:
            self.token_edit.setPlainText(token_value)
        auth_layout.addWidget(self.token_edit)

        token_btn_layout = QHBoxLayout()
        self.save_token_button = QPushButton("Save Token")
        self.save_token_button.clicked.connect(self.save_token)
        token_btn_layout.addStretch()
        token_btn_layout.addWidget(self.save_token_button)
        auth_layout.addLayout(token_btn_layout)

        left_layout.addWidget(auth_group)

        # 2) Store Update group
        store_group = QGroupBox("Stores")
        store_layout = QVBoxLayout(store_group)

        btn_layout = QHBoxLayout()
        self.update_stores_button = QPushButton("Update Stores from API")
        self.update_stores_button.clicked.connect(self.update_stores)
        btn_layout.addWidget(self.update_stores_button)
        btn_layout.addStretch()
        store_layout.addLayout(btn_layout)

        self.store_list_widget = QTreeWidget()
        self.store_list_widget.setHeaderLabels(["Store", "Address"])
        self.store_list_widget.setColumnWidth(0, 400)
        self.store_list_widget.setAlternatingRowColors(True)
        self.store_list_widget.setRootIsDecorated(False)   # hide tree lines
        store_layout.addWidget(self.store_list_widget)

        left_layout.addWidget(store_group)

        # 3) Product Ingestion group
        prod_group = QGroupBox("Product Ingestion")
        prod_layout = QVBoxLayout(prod_group)

        self.product_ids_edit = QPlainTextEdit()
        self.product_ids_edit.setPlaceholderText(
            "e.g.\n202714\n615528\n..."
        )
        self.product_ids_edit.setMaximumBlockCount(10)
        self.product_ids_edit.setFixedHeight(100)
        prod_layout.addWidget(self.product_ids_edit)

        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        prod_layout.addWidget(self.progress_bar)

        btn_hbox = QHBoxLayout()
        self.start_ingestion_button = QPushButton("Start")
        self.start_ingestion_button.clicked.connect(
            self.start_product_ingestion)
        self.start_ingestion_button.setEnabled(False)
        btn_hbox.addWidget(self.start_ingestion_button)

        self.stop_ingestion_button = QPushButton("Stop")
        self.stop_ingestion_button.clicked.connect(self.stop_product_ingestion)
        self.stop_ingestion_button.setEnabled(False)
        btn_hbox.addWidget(self.stop_ingestion_button)
        btn_hbox.addStretch()

        prod_layout.addLayout(btn_hbox)
        left_layout.addWidget(prod_group)

        # 4) Logs
        log_group = QGroupBox("Logs")
        log_layout = QVBoxLayout(log_group)
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        log_layout.addWidget(self.log_text)
        left_layout.addWidget(log_group)

        # ===== RIGHT PANEL =====
        right_panel = QWidget()
        right_layout = QVBoxLayout(right_panel)
        right_layout.setContentsMargins(5, 0, 0, 0)

        # 5) Search Results tree
        results_group = QGroupBox("Stores Having All Products Today (qty > 0)")
        results_layout = QVBoxLayout(results_group)
        self.results_tree = QTreeWidget()
        self.results_tree.setHeaderLabels(["Store / Product", "Details"])
        self.results_tree.header().setStretchLastSection(True)
        self.results_tree.setColumnWidth(0, 300)
        # Alternate row colors
        self.results_tree.setAlternatingRowColors(True)
        results_layout.addWidget(self.results_tree)
        right_layout.addWidget(results_group)

        # Export button
        self.export_button = QPushButton("Export to Excel")
        self.export_button.clicked.connect(self.export_to_excel)
        self.export_button.setEnabled(False)  # enabled only when results exist
        right_layout.addWidget(self.export_button)

        # 6) About Us (styled)
        about_group = QGroupBox("About")
        about_group.setStyleSheet("""
            QGroupBox {
                border: 1px solid #4a86e8;
                margin-top: 1em;
                background-color: #f0f4ff;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 5px 0 5px;
            }
        """)
        about_layout = QVBoxLayout(about_group)
        about_text = (
            '<div style="text-align:center; font-size:8pt; padding:10px;">'
            'Created by <b><a href="https://appuccino.ir">Appuccino Team</a></b> '
            'using <b>DeepSeek</b>!'
            '</div>'
        )
        about_label = QLabel(about_text)
        about_label.setOpenExternalLinks(True)
        about_layout.addWidget(about_label)
        right_layout.addWidget(about_group)

        main_layout.addWidget(left_panel, stretch=1)
        main_layout.addWidget(right_panel, stretch=1)

        # Initial DB load
        self.initial_load_stores()

    # ------------------ helpers & slots ------------------
    def append_log(self, msg):
        self.log_text.append(msg)

    def update_auth_status(self):
        token = self.settings.get("auth_token", "")
        if token:
            self.auth_status_label.setText("✅ Token SET")
            self.auth_status_label.setStyleSheet("color: green;")
        else:
            self.auth_status_label.setText("❌ Token NOT SET")
            self.auth_status_label.setStyleSheet("color: red;")

    def save_token(self):
        token = self.token_edit.toPlainText().strip()
        self.settings["auth_token"] = token
        save_settings(self.settings)
        self.update_auth_status()
        self.append_log("Auth token saved to settings.json")

    def initial_load_stores(self):
        try:
            from db.connection import get_connection
            from repositories.store_repository import get_all_stores_with_info
            conn = get_connection()
            self.store_details = get_all_stores_with_info(conn)
            conn.close()
            self.populate_store_list()
            if self.store_details:
                self.start_ingestion_button.setEnabled(True)
        except Exception as e:
            self.append_log(f"DB load error: {e}")

    def populate_store_list(self):
        self.store_list_widget.clear()
        # self.store_details now contains tuples: (store_id, store_name, location, store_address)
        for sid, name, loc, address in self.store_details:
            item = QTreeWidgetItem(self.store_list_widget)
            item.setText(0, f"{name} (ID: {sid}) - {loc}")
            item.setText(1, address if address else "N/A")
        self.append_log(f"Loaded {len(self.store_details)} stores.")

    def update_stores(self):
        token = self.settings.get("auth_token", "")
        if not token:
            QMessageBox.warning(self, "Missing Token", "Set Auth Token first.")
            return
        try:
            locations = load_store_locations()
        except FileNotFoundError as e:
            QMessageBox.critical(self, "Missing File", str(e))
            return

        cooldown = self.settings.get("cooldown_seconds", 2.5)
        self.update_stores_button.setEnabled(False)
        self.append_log("Starting store update...")

        self.store_worker = StoreUpdateWorker(locations, token, cooldown)
        self.store_worker.log_signal.connect(self.append_log)
        self.store_worker.finished_signal.connect(self.on_stores_updated)
        self.store_worker.start()

    @pyqtSlot(list)
    def on_stores_updated(self, store_details):
        self.store_details = store_details
        self.populate_store_list()
        self.update_stores_button.setEnabled(True)
        self.start_ingestion_button.setEnabled(bool(store_details))

    def parse_product_ids(self):
        raw = self.product_ids_edit.toPlainText()
        parts = raw.replace('\n', ',').split(',')
        return [int(p.strip()) for p in parts if p.strip().isdigit()]

    def start_product_ingestion(self):
        if not self.store_details:
            QMessageBox.information(self, "No Stores", "Update stores first.")
            return
        token = self.settings.get("auth_token", "")
        if not token:
            QMessageBox.warning(self, "Missing Token", "Set Auth Token first.")
            return

        product_ids = self.parse_product_ids()
        if not product_ids:
            QMessageBox.warning(self, "Invalid Input",
                                "Enter valid product IDs.")
            return

        store_ids = [s[0] for s in self.store_details]
        cooldown = self.settings.get("cooldown_seconds", 2.5)

        # Update UI
        self.start_ingestion_button.setEnabled(False)
        self.update_stores_button.setEnabled(False)
        self.stop_ingestion_button.setEnabled(True)
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)

        self.product_worker = ProductIngestionWorker(
            store_ids, product_ids, token, cooldown
        )
        self.product_worker.log_signal.connect(self.append_log)
        self.product_worker.progress_signal.connect(self.update_progress)
        self.product_worker.finished_signal.connect(self.on_ingestion_finished)
        self.product_worker.start()

    def stop_product_ingestion(self):
        if self.product_worker and self.product_worker.isRunning():
            self.product_worker.requestInterruption()
            self.append_log("Stop requested...")
            self.stop_ingestion_button.setEnabled(False)

    def update_progress(self, current, total):
        self.progress_bar.setMaximum(total)
        self.progress_bar.setValue(current)

    @pyqtSlot(list)
    def on_ingestion_finished(self, product_ids):
        self.product_ids_for_search = product_ids
        self.start_ingestion_button.setEnabled(True)
        self.update_stores_button.setEnabled(True)
        self.stop_ingestion_button.setEnabled(False)
        self.progress_bar.setVisible(False)
        self.append_log("Ingestion finished.")
        self.search_and_populate_tree(product_ids)

    def _query_stores_with_all_products(self):
        """Return a list of dicts with store/product data for stores having all products today (qty>0)."""
        product_ids = self.product_ids_for_search
        if not product_ids:
            return []
        try:
            from db.connection import get_connection
            conn = get_connection()
            cursor = conn.cursor()

            placeholders = ','.join(['?'] * len(product_ids))
            # Find stores
            query = f"""
                SELECT ps.store_id
                FROM product_stores ps
                WHERE ps.product_id IN ({placeholders})
                  AND ps.quantity > 0
                  AND date(ps.created_at) = date('now','localtime')
                GROUP BY ps.store_id
                HAVING COUNT(DISTINCT ps.product_id) = ?
            """
            cursor.execute(query, product_ids + [len(product_ids)])
            store_ids = [row[0] for row in cursor.fetchall()]

            if not store_ids:
                conn.close()
                return []

            # Get store info
            cursor.execute(
                f"SELECT store_id, store_name, partner_name, store_address FROM stores WHERE store_id IN ({','.join('?'*len(store_ids))})",
                store_ids
            )
            store_map = {
                row[0]: {
                    'store_id': row[0],
                    'store_name': row[1],
                    'partner_name': row[2],
                    'store_address': row[3] or ''
                }
                for row in cursor.fetchall()
            }

            stores_data = []
            for sid in store_ids:
                cursor.execute(
                    f"""SELECT p.name, ps.quantity, ps.price, ps.ok_price
                        FROM product_stores ps
                        JOIN products p ON p.id = ps.product_id
                        WHERE ps.store_id = ?
                          AND ps.product_id IN ({placeholders})
                          AND ps.quantity > 0
                          AND date(ps.created_at) = date('now','localtime')
                    """,
                    [sid] + product_ids
                )
                rows = cursor.fetchall()
                if not rows:
                    continue

                total_price = sum(row[2] for row in rows)
                total_ok_price = sum(row[3] for row in rows)

                store_info = store_map.get(sid, {})
                store_info['products'] = [
                    {'name': r[0], 'qty': r[1], 'price': r[2], 'ok_price': r[3]} for r in rows]
                store_info['total_price'] = total_price
                store_info['total_ok_price'] = total_ok_price
                stores_data.append(store_info)

            conn.close()
            return stores_data
        except Exception as e:
            self.append_log(f"Error querying data: {e}")
            return []

    def search_and_populate_tree(self, product_ids):
        self.results_tree.clear()
        self.export_button.setEnabled(False)
        if not product_ids:
            return

        stores_data = self._query_stores_with_all_products()
        if not stores_data:
            self.append_log("No store has all products in stock today.")
            return

        def fmt(price):
            """Format price as integer with thousand separators."""
            try:
                return f"{int(price):,}"
            except (ValueError, TypeError):
                return str(price)

        for store in stores_data:
            store_item = QTreeWidgetItem(self.results_tree)
            store_item.setText(0, store['store_name'])
            store_item.setText(
                1, f"ID: {store['store_id']} | Partner: {store.get('partner_name','')}")
            store_item.setExpanded(True)

            for prod in store['products']:
                prod_item = QTreeWidgetItem(store_item)
                prod_item.setText(0, prod['name'])
                prod_item.setText(
                    1, f"Qty: {prod['qty']} | Price: {fmt(prod['price'])} | OK Price: {fmt(prod['ok_price'])}")

            total_item = QTreeWidgetItem(store_item)
            total_item.setText(0, "Total")
            total_item.setText(
                1, f"Sum Price: {fmt(store['total_price'])} | Sum OK Price: {fmt(store['total_ok_price'])}")
            font = total_item.font(0)
            font.setBold(True)
            total_item.setFont(0, font)
            total_item.setFont(1, font)

        self.export_button.setEnabled(True)
        self.append_log("Results tree populated (today's data).")

    def export_to_excel(self):
        if openpyxl is None:
            QMessageBox.critical(self, "Missing Module",
                                 "openpyxl is required. Install it with: pip install openpyxl")
            return

        stores_data = self._query_stores_with_all_products()
        if not stores_data:
            QMessageBox.information(
                self, "No Data", "No stores to export.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel File", "", "Excel Files (*.xlsx)")
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Product Availability"

            # Header (optional)
            ws.append(["Store Name", "Store ID", "Partner", "Address",
                       "Product Name", "Quantity", "Price", "OK Price",
                       "Total Price", "Total OK Price"])

            for store in stores_data:
                # First product row with store totals
                row_data = [
                    store['store_name'],
                    store['store_id'],
                    store.get('partner_name', ''),
                    store.get('store_address', ''),
                    store['products'][0]['name'] if store['products'] else '',
                    store['products'][0]['qty'] if store['products'] else '',
                    store['products'][0]['price'] if store['products'] else '',
                    store['products'][0]['ok_price'] if store['products'] else '',
                    store['total_price'],
                    store['total_ok_price']
                ]
                ws.append(row_data)

                # Subsequent products
                for prod in store['products'][1:]:
                    ws.append(['', '', '', '',
                               prod['name'], prod['qty'], prod['price'], prod['ok_price'],
                               '', ''])

                # Blank row between stores
                ws.append([])

            # Auto-adjust column widths (simple)
            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(
                                max_length, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = min(
                    max_length + 2, 40)

            wb.save(file_path)
            self.append_log(f"Exported to {file_path}")
            QMessageBox.information(
                self, "Success", f"Data exported to:\n{file_path}")
        except Exception as e:
            self.append_log(f"Export error: {e}")
            QMessageBox.critical(self, "Export Error", str(e))
